from pathlib import Path
import json, uuid, re
from docx import Document
from pypdf import PdfReader
from rag.retriever import DATA_DIR
ROOT=Path(__file__).resolve().parents[2]
UPLOAD_DIR=ROOT/'data'/'uploads'; UPLOAD_DIR.mkdir(parents=True,exist_ok=True)
CANDIDATES=DATA_DIR/'legal_documents_candidates.json'
APPROVED=DATA_DIR/'legal_documents_approved.json'

def _looks_corrupted(text):
    if not text or not text.strip():
        return True
    # Common indicators of broken PDF ToUnicode maps / mojibake.
    bad = ('\ufffd', '\x00', 'Ã', 'Â', 'â€™', 'â€œ', 'â€', 'ðŸ')
    if any(token in text for token in bad):
        return True
    visible = [c for c in text if not c.isspace()]
    if not visible:
        return True
    odd = sum(1 for c in visible if (ord(c) < 32 or 0x7f <= ord(c) <= 0x9f))
    return odd / max(1, len(visible)) > 0.01


def _ocr_pdf(path, language='eng+kan+tam+tel+mal+hin'):
    try:
        import pymupdf
        doc = pymupdf.open(str(path))
        pages = []
        for page in doc:
            # OCR every page only when native text is empty or appears corrupted.
            native = page.get_text('text') or ''
            if len(native.strip()) >= 40 and not _looks_corrupted(native):
                pages.append(native)
                continue
            try:
                textpage = page.get_textpage_ocr(language=language, dpi=250, full=True)
                pages.append(page.get_text('text', textpage=textpage) or '')
            except Exception as exc:
                raise ValueError(
                    'PDF text appears scanned or corrupted and OCR could not run. Install Tesseract OCR and '
                    'language data for English/Kannada/Tamil/Telugu/Malayalam/Hindi, then retry. '
                    'If the PDF has a broken font map, OCR is needed instead of native text extraction. '
                    f'OCR detail: {exc}'
                ) from exc
        return '\n'.join(pages)
    except ImportError as exc:
        raise ValueError('PDF text appears scanned or corrupted. Install PyMuPDF and Tesseract OCR with the required language data to recover it.') from exc


def extract_text(path, ocr_languages='eng+kan+tam+tel+mal+hin'):
    ext=path.suffix.lower()
    if ext=='.pdf':
        text='\n'.join(p.extract_text() or '' for p in PdfReader(str(path)).pages)
        if len(text.strip()) < 40 or _looks_corrupted(text):
            ocr_text = _ocr_pdf(path, ocr_languages)
            if len(ocr_text.strip()) >= 20 and not _looks_corrupted(ocr_text):
                return ocr_text
            if len(text.strip()) < 40:
                raise ValueError('No readable text could be extracted. This PDF may be scanned, protected, or missing OCR language data. Try a clearer/searchable PDF or install the appropriate Tesseract language packs.')
            raise ValueError('The PDF text encoding appears corrupted, and OCR did not produce reliable text. Please use a clearer PDF or scan it at 300 DPI and ensure the correct language pack is installed.')
        return text
    if ext=='.docx': return '\n'.join(p.text for p in Document(str(path)).paragraphs)
    if ext in {'.txt','.md'}: return path.read_text(encoding='utf-8',errors='replace')
    raise ValueError('Supported file types: PDF, DOCX, TXT, MD')

def chunk(text,size=1300,overlap=180):
    text=re.sub(r'[ \t]+',' ',text).strip(); out=[]; start=0
    while start<len(text):
        end=min(len(text),start+size)
        if end<len(text):
            boundary=max(text.rfind('. ',start,end),text.rfind('\n',start,end))
            if boundary>start+size//2:end=boundary+1
        part=text[start:end].strip()
        if len(part)>20: out.append(part)
        if end>=len(text):break
        start=max(start+1,end-overlap)
    return out

def load(path):
    if not path.exists():return []
    try:return json.loads(path.read_text(encoding='utf-8'))
    except Exception:return []

def ingest(path,title,source_url,jurisdiction,effective_date,category,review_status='pending_review'):
    text=extract_text(path)
    if len(text.strip())<30: raise ValueError('No readable text found. Scanned PDFs need OCR first.')
    chunks=chunk(text); sid='USER-'+uuid.uuid4().hex[:12].upper()
    records=[]
    for i,content in enumerate(chunks):
        records.append({'chunk_id':f'{sid}-C{i+1:04d}','source_id':sid,'title':title or path.name,
          'authority':jurisdiction or 'Not specified','jurisdiction':jurisdiction or 'Not specified','category':category or 'general',
          'content':content,'provenance':'user_uploaded_document','review_status':review_status,
          'official_source_url':source_url or None,'source_url':source_url or None,'effective_date':effective_date or 'not_verified',
          'document_filename':path.name,'notice':'Uploaded source; authority/current validity not independently verified.'})
    # User documents are searchable only after review/promotion, but stored as candidates by default.
    arr=load(CANDIDATES); arr.extend(records); CANDIDATES.write_text(json.dumps(arr,ensure_ascii=False,indent=2),encoding='utf-8')
    return {'source_id':sid,'title':title or path.name,'chunks':len(records),'review_status':review_status,
            'message':'Stored as pending review. Review provenance, jurisdiction and current validity before promoting into RAG.'}

def promote(source_id):
    candidates=load(CANDIDATES); approved=load(APPROVED); selected=[x for x in candidates if x.get('source_id')==source_id]
    if not selected:return False
    for x in selected:x['review_status']='approved'
    approved.extend(selected); APPROVED.write_text(json.dumps(approved,ensure_ascii=False,indent=2),encoding='utf-8')
    remaining=[x for x in candidates if x.get('source_id')!=source_id]; CANDIDATES.write_text(json.dumps(remaining,ensure_ascii=False,indent=2),encoding='utf-8')
    return True

def list_sources():
    rows=[]
    for path,status in [(APPROVED,'approved'),(CANDIDATES,'pending_review')]:
        grouped={}
        for x in load(path):
            sid=x.get('source_id','unknown'); grouped.setdefault(sid,{'source_id':sid,'title':x.get('title'),'jurisdiction':x.get('jurisdiction'),'source_url':x.get('source_url'),'effective_date':x.get('effective_date'),'category':x.get('category'),'status':status,'chunks':0})['chunks']+=1
        rows.extend(grouped.values())
    return rows

def make_template(template_id,outdir):
    outdir.mkdir(parents=True,exist_ok=True)
    if template_id=='evidence_checklist':
        p=outdir/'NyayaSetu_Evidence_Checklist.txt'
        p.write_text('NYAYASETU — EVIDENCE CHECKLIST\n\n[ ] Timeline with dates\n[ ] Relevant contract/agreement/notice/order\n[ ] Payment/salary/transaction proof where relevant\n[ ] Emails, messages and written requests\n[ ] Replies and proof of delivery\n[ ] Witness details, if appropriate\n[ ] Deadline or hearing date noted\n[ ] Copies backed up securely\n[ ] Unnecessary personal identifiers redacted\n\nThis is a general checklist, not an official form or legal advice. Confirm the requirements with the relevant authority or a qualified advocate.\n',encoding='utf-8'); return p
    names={'representation':'General Representation / Complaint — Editable Starter','notice_reply':'Reply to Notice — Editable Starter'}
    if template_id not in names:raise KeyError(template_id)
    p=outdir/('NyayaSetu_Representation_Template.docx' if template_id=='representation' else 'NyayaSetu_Notice_Reply_Template.docx')
    d=Document(); d.add_heading(names[template_id],0)
    d.add_paragraph('IMPORTANT: Generic editable drafting aid, not an official government form or legal advice. Confirm the prescribed format, forum, deadline, fee and filing method with the relevant authority or a qualified advocate.')
    sections=[('1. Recipient / Authority','[Name and designation]\n[Office address / official email]'),('2. Applicant / Sender','Name: [Full name]\nAddress / contact: [Only details required for filing]'),('3. Subject','[One-line description of the issue and requested action]'),('4. Facts in date order','1. [Date] — [Factual event and supporting record]\n2. [Date] — [Factual event and supporting record]\n3. [Date] — [Factual event and supporting record]'),('5. Previous steps taken','[Requests/complaints made, dates, proof of delivery and replies]'),('6. Requested resolution','[State the specific action requested]'),('7. Documents attached','1. [Document name and date]\n2. [Document name and date]\n3. [Other evidence]'),('8. Declaration and signature','I believe the factual statements above are accurate to the best of my knowledge.\n\nPlace: [City]    Date: [DD/MM/YYYY]\nSignature: [Name]'),('Before sending','Verify the correct authority, current law, prescribed form, filing channel, fee and deadline. Keep a complete copy and proof of submission.')]
    for h,b in sections:d.add_heading(h,level=1);d.add_paragraph(b)
    d.save(p);return p
